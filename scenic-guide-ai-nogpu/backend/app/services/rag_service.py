"""
RAG知识库服务 - 核心AI问答引擎
"""
import asyncio
import csv
import io
import os
import uuid
from typing import List, Optional, Dict, Any
from loguru import logger

import chromadb
from chromadb.config import Settings as ChromaSettings
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    PyPDFLoader, Docx2txtLoader, TextLoader, UnstructuredMarkdownLoader
)
from langchain_core.documents import Document

from app.core.config import settings
from openai import AsyncOpenAI

# 共享embedding客户端
_embedding_client: Optional[AsyncOpenAI] = None


def _get_embedding_client() -> AsyncOpenAI:
    global _embedding_client
    if _embedding_client is None:
        _embedding_client = AsyncOpenAI(
            api_key=settings.OPENAI_API_KEY,
            base_url="https://ark.cn-beijing.volces.com/api/v3",
        )
    return _embedding_client


def _extract_xlsx(file_path: str) -> str:
    """从 .xlsx 文件提取纯文本"""
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    texts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_lines = [f"【工作表: {sheet_name}】"]
        for row in ws.iter_rows(values_only=True):
            row_text = " | ".join([str(c) if c is not None else "" for c in row]).strip()
            if row_text:
                sheet_lines.append(row_text)
        texts.append("\n".join(sheet_lines))
    wb.close()
    return "\n\n".join(texts)


def _extract_xls(file_path: str) -> str:
    """从 .xls 文件提取纯文本"""
    try:
        import xlrd
        wb = xlrd.open_workbook(file_path)
        texts = []
        for sheet in wb.sheets():
            sheet_lines = [f"【工作表: {sheet.name}】"]
            for row_idx in range(sheet.nrows):
                row_values = sheet.row_values(row_idx)
                row_text = " | ".join([str(c) if c != "" else "" for c in row_values]).strip()
                if row_text:
                    sheet_lines.append(row_text)
            texts.append("\n".join(sheet_lines))
        return "\n\n".join(texts)
    except ImportError:
        raise RuntimeError("需要安装 xlrd 库来读取 .xls 文件: pip install xlrd")


def _extract_csv(file_path: str) -> str:
    """从 .csv 文件提取纯文本"""
    texts = []
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        reader = csv.reader(f)
        for row in reader:
            row_text = " | ".join(row).strip()
            if row_text:
                texts.append(row_text)
    return "\n".join(texts)


def extract_file_text(file_path: str) -> str:
    """根据文件扩展名提取文本内容（用于预览等，不入库）"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        loader = PyPDFLoader(file_path)
        docs = loader.load()
        return "\n\n".join([d.page_content for d in docs])
    elif ext in (".docx", ".doc"):
        loader = Docx2txtLoader(file_path)
        docs = loader.load()
        return "\n\n".join([d.page_content for d in docs])
    elif ext == ".md":
        loader = UnstructuredMarkdownLoader(file_path)
        docs = loader.load()
        return "\n\n".join([d.page_content for d in docs])
    elif ext == ".xlsx":
        return _extract_xlsx(file_path)
    elif ext == ".xls":
        return _extract_xls(file_path)
    elif ext == ".csv":
        return _extract_csv(file_path)
    else:
        loader = TextLoader(file_path, encoding="utf-8")
        docs = loader.load()
        return "\n\n".join([d.page_content for d in docs])


class RAGService:
    """RAG检索增强生成服务"""

    def __init__(self):
        self.client = None
        self.collection = None
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.RAG_CHUNK_SIZE,
            chunk_overlap=settings.RAG_CHUNK_OVERLAP,
            separators=["\n\n", "\n", "。", "！", "？", "；", " ", ""],
        )
        self._embedding_fn = None
        self._init_chroma()

    def _init_chroma(self):
        """初始化ChromaDB"""
        try:
            os.makedirs(settings.CHROMA_PERSIST_DIR, exist_ok=True)
            self.client = chromadb.PersistentClient(
                path=settings.CHROMA_PERSIST_DIR,
                settings=ChromaSettings(anonymized_telemetry=False)
            )
            self.collection = self.client.get_or_create_collection(
                name=settings.CHROMA_COLLECTION,
                metadata={"hnsw:space": "cosine"}
            )
            logger.info(f"ChromaDB初始化成功，集合: {settings.CHROMA_COLLECTION}，"
                       f"文档数: {self.collection.count()}")
        except Exception as e:
            logger.error(f"ChromaDB初始化失败: {e}")

    async def _get_embedding_async(self, text: str) -> List[float]:
        """获取文本嵌入向量（使用火山引擎embedding API）"""
        try:
            client = _get_embedding_client()
            resp = await client.embeddings.create(
                model=settings.EMBEDDING_MODEL,
                input=text
            )
            return resp.data[0].embedding
        except Exception as e:
            logger.warning(f"Embedding API失败，使用随机向量: {e}")
            import random
            return [random.random() for _ in range(1024)]

    async def add_document(
        self,
        content: str,
        metadata: Dict[str, Any],
        doc_id: Optional[str] = None
    ) -> int:
        """添加文档到知识库，返回切片数量"""
        try:
            # 文本切片
            documents = self.text_splitter.create_documents(
                texts=[content],
                metadatas=[metadata]
            )
            
            chunk_ids = []
            chunk_texts = []
            chunk_embeddings = []
            chunk_metadatas = []
            
            base_id = doc_id or str(uuid.uuid4())
            
            for i, doc in enumerate(documents):
                chunk_id = f"{base_id}_chunk_{i}"
                chunk_ids.append(chunk_id)
                chunk_texts.append(doc.page_content)
                # 直接 await 异步 embedding 函数（asyncio.to_thread 不能用于协程）
                embedding = await self._get_embedding_async(doc.page_content)
                chunk_embeddings.append(embedding)
                meta = doc.metadata.copy()
                meta["chunk_index"] = i
                meta["total_chunks"] = len(documents)
                chunk_metadatas.append(meta)
            
            if chunk_ids:
                self.collection.add(
                    ids=chunk_ids,
                    documents=chunk_texts,
                    embeddings=chunk_embeddings,
                    metadatas=chunk_metadatas
                )
            
            logger.info(f"文档 {base_id} 已切分为 {len(documents)} 个片段并入库")
            return len(documents)
            
        except Exception as e:
            logger.error(f"添加文档失败: {e}")
            raise

    async def load_file(self, file_path: str, metadata: Dict[str, Any]) -> int:
        """加载文件到知识库"""
        ext = os.path.splitext(file_path)[1].lower()

        try:
            if ext == ".pdf":
                loader = PyPDFLoader(file_path)
                docs = loader.load()
                full_content = "\n\n".join([d.page_content for d in docs])
            elif ext in (".docx", ".doc"):
                loader = Docx2txtLoader(file_path)
                docs = loader.load()
                full_content = "\n\n".join([d.page_content for d in docs])
            elif ext == ".md":
                loader = UnstructuredMarkdownLoader(file_path)
                docs = loader.load()
                full_content = "\n\n".join([d.page_content for d in docs])
            elif ext == ".xlsx":
                full_content = _extract_xlsx(file_path)
            elif ext == ".xls":
                full_content = _extract_xls(file_path)
            elif ext == ".csv":
                full_content = _extract_csv(file_path)
            else:  # .txt 及其他
                loader = TextLoader(file_path, encoding="utf-8")
                docs = loader.load()
                full_content = "\n\n".join([d.page_content for d in docs])

            # 检查内容长度，防止超大文件撑爆内存和API配额
            max_chars = settings.MAX_CONTENT_CHARS
            if len(full_content) > max_chars:
                raise ValueError(
                    f"文档内容过大（{len(full_content):,} 字符），"
                    f"最大支持 {max_chars:,} 字符（约 {max_chars // settings.RAG_CHUNK_SIZE} 个分块）。"
                    f"请拆分后分批上传。"
                )

            return await self.add_document(
                content=full_content,
                metadata=metadata,
                doc_id=metadata.get("doc_id")
            )
        except Exception as e:
            logger.error(f"加载文件失败 {file_path}: {e}")
            raise

    async def search(
        self,
        query: str,
        top_k: int = None,
        filter_metadata: Optional[Dict] = None
    ) -> List[Dict[str, Any]]:
        """语义搜索，返回相关文档片段"""
        top_k = top_k or settings.RAG_TOP_K
        
        try:
            # 直接await异步版本（search是async函数，避免死锁）
            query_embedding = await self._get_embedding_async(query)
            
            where_clause = filter_metadata if filter_metadata else None
            
            results = self.collection.query(
                query_embeddings=[query_embedding],
                n_results=min(top_k, max(1, self.collection.count())),
                where=where_clause,
                include=["documents", "metadatas", "distances"]
            )
            
            retrieved = []
            if results and results["documents"]:
                for doc, meta, dist in zip(
                    results["documents"][0],
                    results["metadatas"][0],
                    results["distances"][0]
                ):
                    similarity = 1 - dist  # cosine distance → similarity
                    if similarity >= settings.RAG_SIMILARITY_THRESHOLD:
                        retrieved.append({
                            "content": doc,
                            "metadata": meta,
                            "similarity": round(similarity, 4)
                        })
            
            logger.debug(f"RAG搜索 '{query[:30]}...' 返回 {len(retrieved)} 条结果")
            return retrieved
            
        except Exception as e:
            logger.error(f"RAG搜索失败: {e}")
            return []

    async def delete_document(self, doc_id: str) -> bool:
        """删除文档的所有切片"""
        try:
            # 查找所有相关chunk
            results = self.collection.get(
                where={"doc_id": doc_id}
            )
            if results["ids"]:
                self.collection.delete(ids=results["ids"])
                logger.info(f"已删除文档 {doc_id} 的 {len(results['ids'])} 个切片")
            return True
        except Exception as e:
            logger.error(f"删除文档失败: {e}")
            return False

    async def get_document_chunks(self, doc_id: str) -> List[Dict[str, Any]]:
        """获取文档的所有切片（按 chunk_index 排序）"""
        try:
            results = self.collection.get(
                where={"doc_id": doc_id},
                include=["documents", "metadatas"]
            )
            chunks = []
            if results and results["ids"]:
                for i, chunk_id in enumerate(results["ids"]):
                    chunks.append({
                        "id": chunk_id,
                        "content": results["documents"][i] if results["documents"] else "",
                        "metadata": results["metadatas"][i] if results["metadatas"] else {},
                    })
                # 按 chunk_index 排序
                chunks.sort(key=lambda c: c["metadata"].get("chunk_index", 0))
            return chunks
        except Exception as e:
            logger.error(f"获取文档切片失败: {e}")
            return []

    def get_stats(self) -> Dict[str, Any]:
        """获取知识库统计信息"""
        try:
            count = self.collection.count()
            return {"total_chunks": count, "collection": settings.CHROMA_COLLECTION}
        except Exception as e:
            return {"total_chunks": 0, "error": str(e)}


# 全局RAG服务实例
rag_service = RAGService()
